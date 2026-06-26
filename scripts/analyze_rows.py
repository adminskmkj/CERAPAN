import openpyxl
from openpyxl.utils import get_column_letter

path = r"C:\Users\user\Downloads\08-Tapak_Standard_4 PdP Updated 28.10.2024 (2) (1).xlsm"
wb = openpyxl.load_workbook(path, data_only=False)
ws = wb["STANDARD 4"]

for r in range(8, 20):
    parts = []
    for c in range(2, 20):
        v = ws.cell(r, c).value
        if v:
            text = str(v).replace("\n", " | ")[:80]
            parts.append(f"{get_column_letter(c)}={text}")
    if parts:
        print(f"R{r}: " + " || ".join(parts))