import openpyxl

path = r"C:\Users\user\Downloads\08-Tapak_Standard_4 PdP Updated 28.10.2024 (2) (1).xlsm"
wb = openpyxl.load_workbook(path, data_only=False, keep_vba=True)
ws = wb["STANDARD 4"]

print("=== Column B - Tindakan descriptions ===")
for r in range(8, 20):
    v = ws.cell(r, 2).value
    if v:
        print(f"B{r}: {str(v).replace(chr(10), ' | ')[:180]}")

for start, label in [(30, "4.2"), (55, "4.2.2"), (82, "4.3"), (117, "4.4"), (163, "4.4.2"), (190, "4.5"), (230, "4.6")]:
    print(f"\n--- {label} section ---")
    for r in range(start, start + 12):
        v = ws.cell(r, 2).value
        if v:
            print(f"B{r}: {str(v).replace(chr(10), ' | ')[:180]}")

print("\n=== Row 4 Arahan ===")
print(ws.cell(4, 2).value)